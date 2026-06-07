"""
commission_export.py — Fill CG Line+ or Easicheck commissioning calculator,
                       convert to PDF via LibreOffice headless.
"""
import sys, os, argparse, subprocess


def find_libreoffice():
    import shutil as sh
    for cmd in ["soffice", "libreoffice"]:
        p = sh.which(cmd)
        if p:
            return p
    for p in [
        r"C:\Users\E0740516\Downloads\LibreOfficePortable\App\libreoffice\program\soffice.exe",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]:
        if os.path.exists(p):
            return p
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--type",           required=True, choices=["cg", "easicheck"])
    ap.add_argument("--panels",         required=True, type=int)
    ap.add_argument("--lumis",          required=True, type=int)
    ap.add_argument("--cards",          required=True, type=int)
    ap.add_argument("--software",       required=True, choices=["yes", "no"])
    ap.add_argument("--central-london", required=True, choices=["yes", "no"], dest="central_london")
    ap.add_argument("--outdir",         required=True)
    args = ap.parse_args()

    here     = os.path.dirname(os.path.abspath(__file__))
    template = os.path.join(here, "docs", "commission_calculators.xlsx")

    if not os.path.exists(template):
        print(f"__ERROR__:commission_calculators.xlsx not found at {template}")
        sys.exit(1)

    lo = find_libreoffice()
    if not lo:
        print("__ERROR__:LibreOffice not found")
        sys.exit(1)

    os.makedirs(args.outdir, exist_ok=True)
    work_file = os.path.join(args.outdir, "Commission_Calculation.xlsx")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(template, keep_vba=False)

        panels   = args.panels
        lumis    = args.lumis
        cards    = args.cards
        sw_qty   = 1 if args.software == "yes" else 0
        cl_val   = "Yes" if args.central_london == "yes" else "No"

        if args.type == "cg":
            ws = wb["CG Line +"]
            sheet_name = "CG Line +"

            # Zero all qty input cells
            for cell in ("C5","C6","C7","C12","C13","C14","C19","C20","C21","C26","C27","C28"):
                ws[cell] = 0

            # Fill the correct tier row based on panel count
            if panels == 1:
                ws["C5"] = panels; ws["C6"] = lumis; ws["C7"] = sw_qty
            elif panels <= 4:
                ws["C12"] = panels; ws["C13"] = lumis; ws["C14"] = sw_qty
            elif panels <= 9:
                ws["C19"] = panels; ws["C20"] = lumis; ws["C21"] = sw_qty
            elif panels <= 15:
                ws["C26"] = panels; ws["C27"] = lumis; ws["C28"] = sw_qty
            else:
                print("__ERROR__:CG Line+ supports up to 15 panels only")
                sys.exit(1)

            ws["K12"] = cl_val

        else:  # easicheck
            ws = wb["Easicheck"]
            sheet_name = "Easicheck"

            # Zero all qty input cells
            for cell in ("C5","C6","C7",
                         "C11","C12","C13","C14",
                         "C18","C19","C20","C21",
                         "C25","C26","C27","C28",
                         "C32","C33","C34","C35"):
                ws[cell] = 0

            # Fill the correct tier row
            if panels == 1:
                ws["C5"] = panels; ws["C6"] = lumis; ws["C7"] = sw_qty
            elif panels <= 4:
                ws["C11"] = panels; ws["C12"] = lumis; ws["C13"] = cards; ws["C14"] = sw_qty
            elif panels <= 9:
                ws["C18"] = panels; ws["C19"] = lumis; ws["C20"] = cards; ws["C21"] = sw_qty
            elif panels <= 15:
                ws["C25"] = panels; ws["C26"] = lumis; ws["C27"] = cards; ws["C28"] = sw_qty
            else:
                ws["C32"] = panels; ws["C33"] = lumis; ws["C34"] = cards; ws["C35"] = sw_qty

            ws["K12"] = cl_val

        # Hide all other sheets — LibreOffice skips hidden sheets in PDF export
        for name in wb.sheetnames:
            if name != sheet_name:
                wb[name].sheet_state = "hidden"
        wb.active = wb[sheet_name]

        wb.calculation.fullCalcOnLoad = True
        wb.save(work_file)

    except Exception as e:
        print(f"__ERROR__:Failed to fill template: {e}")
        sys.exit(1)

    # ── Convert to PDF via LibreOffice ────────────────────────────────────────
    try:
        result = subprocess.run(
            [lo, "--headless", "--convert-to", "pdf", "--outdir", args.outdir, work_file],
            capture_output=True, text=True, timeout=90,
        )
        if result.returncode != 0:
            print(f"__ERROR__:LibreOffice failed: {result.stderr.strip()}")
            sys.exit(1)
    except subprocess.TimeoutExpired:
        print("__ERROR__:LibreOffice timed out after 90 s")
        sys.exit(1)
    except Exception as e:
        print(f"__ERROR__:LibreOffice error: {e}")
        sys.exit(1)

    pdf = os.path.join(args.outdir, "Commission_Calculation.pdf")
    if not os.path.exists(pdf):
        alt = work_file.replace(".xlsx", ".pdf")
        if os.path.exists(alt):
            os.rename(alt, pdf)
        else:
            print("__ERROR__:PDF not found after conversion")
            sys.exit(1)

    print(f"__PDF__:{pdf}")
    sys.exit(0)


if __name__ == "__main__":
    main()
