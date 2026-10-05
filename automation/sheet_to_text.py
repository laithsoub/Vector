"""
sheet_to_text.py — Vector
Turns a spreadsheet attachment (.xlsx/.xlsm/.xlsb/.xls/.csv) into plain text
the AI can read. Gemini takes no Excel MIME type, so the Inbox Summarize / chat
used to drop spreadsheets silently; this renders each sheet as tab-separated
rows (cached VALUES, not formulas) with a hard size cap.

Usage: sheet_to_text.py --input FILE [--max-chars N]
Prints JSON: {"text": "..."} or {"error": "..."}
"""

import argparse
import csv
import json
import os
import sys

MAX_ROWS_PER_SHEET = 400
MAX_COLS = 40
MAX_SCAN = 20000   # rows walked per sheet just to report the total


def _fmt(v) -> str:
    if v is None:
        return ''
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f'{v:.6g}'
    if hasattr(v, 'strftime'):
        return v.strftime('%Y-%m-%d')
    return str(v).replace('\t', ' ').replace('\r', ' ').replace('\n', ' ').strip()


def _render(name: str, rows) -> str:
    """Tab-separated, empty rows dropped, trailing empty cells trimmed."""
    out, kept, total = [], 0, 0
    for r in rows:
        cells = [_fmt(v) for v in list(r)[:MAX_COLS]]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            continue
        total += 1
        if total > MAX_SCAN:
            break
        if kept < MAX_ROWS_PER_SHEET:
            out.append('\t'.join(cells))
            kept += 1
    if not out:
        return ''
    head = f'### Sheet: {name}'
    if total > kept:
        head += f' (first {kept} of {"20000+" if total > MAX_SCAN else total} non-empty rows)'
    return head + '\n' + '\n'.join(out)


def _xlsx(path: str):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if getattr(ws, 'sheet_state', 'visible') != 'visible':
                continue
            yield _render(ws.title, ws.iter_rows(values_only=True))
    finally:
        wb.close()


def _xlsb(path: str):
    from pyxlsb import open_workbook
    with open_workbook(path) as wb:
        for name in wb.sheets:
            with wb.get_sheet(name) as ws:
                yield _render(name, ([c.v for c in row] for row in ws.rows()))


def _csv(path: str):
    with open(path, newline='', encoding='utf-8-sig', errors='replace') as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        yield _render(os.path.basename(path), csv.reader(f, dialect))


def _xls(path: str):
    # Legacy .xls: no xlrd on these machines, so ask Excel itself (COM).
    import win32com.client
    xl = win32com.client.DispatchEx('Excel.Application')
    xl.Visible = False
    xl.DisplayAlerts = False
    wb = None
    try:
        wb = xl.Workbooks.Open(os.path.abspath(path), ReadOnly=True, UpdateLinks=0)
        for ws in wb.Worksheets:
            if ws.Visible != -1:
                continue
            vals = ws.UsedRange.Value
            if vals is None:
                continue
            if not isinstance(vals, tuple):
                vals = ((vals,),)
            yield _render(ws.Name, vals)
    finally:
        if wb is not None:
            wb.Close(False)
        xl.Quit()


READERS = {'.xlsx': _xlsx, '.xlsm': _xlsx, '.xlsb': _xlsb, '.csv': _csv, '.xls': _xls}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--input', required=True)
    ap.add_argument('--max-chars', type=int, default=60000)
    a = ap.parse_args()

    reader = READERS.get(os.path.splitext(a.input)[1].lower())
    if not reader:
        print(json.dumps({'error': 'not a spreadsheet'}))
        return
    try:
        text = '\n\n'.join(s for s in reader(a.input) if s)
    except Exception as e:
        print(json.dumps({'error': f'could not read spreadsheet: {e}'}))
        return
    if len(text) > a.max_chars:
        text = text[:a.max_chars] + '\n… (truncated)'
    print(json.dumps({'text': text}, ensure_ascii=False))


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
