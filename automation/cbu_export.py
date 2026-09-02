"""
cbu_export.py — Fill CBU Tech Brief, convert to PDF, then append Commissioning
                and T&C PDFs from the docs/ folder (T&C always last).
"""
import sys, os, argparse, subprocess


def find_libreoffice():
    import shutil as sh
    for cmd in ["soffice", "libreoffice"]:
        p = sh.which(cmd)
        if p:
            return p
    for p in [
        r"C:\Users\E0740516\Downloads\LibreOfficePortable\App\libreoffice\program\soffice.exe",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]:
        if os.path.exists(p):
            return p
    return None


def find_doc(docs_dir, *candidates):
    for name in candidates:
        p = os.path.join(docs_dir, name)
        if os.path.exists(p):
            return p
    return None


def merge_pdfs(output_path, paths):
    """Merge PDFs using pypdf (falls back to PyPDF2).  Returns (ok, errmsg)."""
    try:
        from pypdf import PdfWriter, PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfWriter, PdfReader
        except ImportError:
            return False, "pypdf not installed — brief exported without appendices"

    writer = PdfWriter()
    for p in paths:
        if not p or not os.path.exists(p):
            continue
        try:
            for page in PdfReader(p).pages:
                writer.add_page(page)
        except Exception as e:
            sys.stderr.write(f"[warn] could not read {p}: {e}\n")

    with open(output_path, "wb") as f:
        writer.write(f)
    return True, ""


def patch_sheet_errors(wb):
    """Correct the known-wrong rows of the 'Tables' sheet before the brief is
    rendered from it.

    The calculator is maintained by hand and carries errors from revision to
    revision. 1PH-12KVA is a 3x 5KVA parallel build (the brief itself prints 3
    control cabinets), but its supply row still holds the 2x fuse and cable
    figures copied down from 1PH-10KVA — still wrong in V3.6. Every other
    multi-cabinet row follows the Nx rule, and automation/cbu_data_gen.py applies
    the same correction to src/lib/cbuData.ts, so without this the printed brief
    contradicted the app for the same system.

    Only rewrites cells that still hold the stale value, so the day the sheet is
    fixed upstream this quietly stops doing anything.
    """
    if "Tables" not in wb.sheetnames:
        return
    ws = wb["Tables"]
    for row in ws.iter_rows(min_row=22, max_row=54, max_col=12):
        if str(row[0].value or "").strip() != "1PH- 12KVA":
            continue
        fuse_rect, fuse_bypass, cable_in, cable_out = row[8], row[9], row[10], row[11]
        for cell in (fuse_rect, fuse_bypass):
            if str(cell.value or "").strip() == "2x40A":
                cell.value = "3x40A"
        for cell in (cable_in, cable_out):
            if str(cell.value or "").strip().startswith("2x"):
                cell.value = "3x 10mm²"
        break


def fill_and_convert(template, lo, outdir, idx, system, project, quote, engineer, email, phone):
    """Fill template for one system, convert to PDF.  Returns path to PDF or raises."""
    import openpyxl
    from openpyxl.worksheet.page import PrintPageSetup
    from openpyxl.worksheet.properties import WorksheetProperties, PageSetupProperties

    work_file = os.path.join(outdir, f"CBU_Tech_Brief_{idx}.xlsx")

    wb = openpyxl.load_workbook(template, keep_vba=True)
    patch_sheet_errors(wb)

    if "Sales Engineer Sheet" in wb.sheetnames:
        ws = wb["Sales Engineer Sheet"]
        ws["B5"]  = system
        ws["C17"] = project
        ws["C18"] = quote
        ws["C19"] = engineer

    if "CBU Tech Brief" in wb.sheetnames:
        ws2 = wb["CBU Tech Brief"]
        ws2["F5"] = project
        ws2["F6"] = quote
        ws2["F7"] = engineer
        ws2["F8"] = email
        ws2["F9"] = phone
        ws2.page_setup = PrintPageSetup(
            paperSize=9, orientation="portrait",
            fitToWidth=1, fitToHeight=0,
        )
        if ws2.sheet_properties is None:
            ws2.sheet_properties = WorksheetProperties()
        ws2.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    # HIDE every other sheet — do NOT delete them.
    # CBU Tech Brief has XLOOKUP formulas that reference Sales Engineer Sheet;
    # deleting it makes every lookup produce Err:504 before LibreOffice can calc.
    for name in wb.sheetnames:
        if name != "CBU Tech Brief":
            wb[name].sheet_state = "hidden"

    wb.active = wb["CBU Tech Brief"]
    wb.calculation.fullCalcOnLoad = True
    wb.save(work_file)

    result = subprocess.run(
        [lo, "--headless", "--convert-to", "pdf", "--outdir", outdir, work_file],
        capture_output=True, text=True, timeout=90,
    )
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice failed: {result.stderr.strip()}")

    pdf = os.path.join(outdir, f"CBU_Tech_Brief_{idx}.pdf")
    if not os.path.exists(pdf):
        alt = work_file.replace(".xlsx", ".pdf")
        if os.path.exists(alt):
            os.rename(alt, pdf)
        else:
            raise RuntimeError(f"PDF not found after conversion (system {idx}: {system})")
    return pdf


def main():
    ap = argparse.ArgumentParser()
    # --system may be specified multiple times for multi-system exports
    ap.add_argument("--system",   action="append", dest="systems", required=True)
    ap.add_argument("--project",  required=True)
    ap.add_argument("--quote",    required=True)
    ap.add_argument("--engineer", required=True)
    ap.add_argument("--email",    required=True)
    ap.add_argument("--phone",    required=True)
    ap.add_argument("--outdir",   required=True)
    args = ap.parse_args()

    here     = os.path.dirname(os.path.abspath(__file__))
    template = os.path.join(here, "cbu_calculator.xlsm")
    docs_dir = os.path.join(here, "docs")

    if not os.path.exists(template):
        print(f"__ERROR__:cbu_calculator.xlsm is missing from this install ({template})")
        sys.exit(1)

    lo = find_libreoffice()
    if not lo:
        print("__ERROR__:LibreOffice is not installed — it is what converts the sheet to PDF")
        sys.exit(1)

    os.makedirs(args.outdir, exist_ok=True)

    # ── Generate one brief PDF per system ─────────────────────────────────────
    brief_pdfs = []
    for idx, system in enumerate(args.systems, 1):
        try:
            pdf = fill_and_convert(
                template, lo, args.outdir, idx, system,
                args.project, args.quote, args.engineer, args.email, args.phone,
            )
            brief_pdfs.append(pdf)
        except subprocess.TimeoutExpired:
            print(f"__ERROR__:LibreOffice timed out on system {idx} ({system})")
            sys.exit(1)
        except Exception as e:
            print(f"__ERROR__:{e}")
            sys.exit(1)

    # ── Append Commissioning + T&C once at end ────────────────────────────────
    comm_pdf = find_doc(docs_dir,
        "commissioning.pdf",
        "GENERAL SERVICE & COMMISSIONING.pdf",
    )
    tc_pdf = find_doc(docs_dir,
        "terms_and_conditions.pdf",
        "United Kingdom TCs English March 2023.pdf",
    )

    final_pdf = brief_pdfs[0] if len(brief_pdfs) == 1 else None
    to_merge = brief_pdfs[:]
    if comm_pdf:
        to_merge.append(comm_pdf)
    if tc_pdf:
        to_merge.append(tc_pdf)   # T&C always last

    if len(to_merge) > 1:
        merged = os.path.join(args.outdir, "CBU_Brief_Complete.pdf")
        ok, err = merge_pdfs(merged, to_merge)
        if ok and os.path.exists(merged):
            final_pdf = merged
        else:
            sys.stderr.write(f"[warn] PDF merge skipped: {err}\n")
            final_pdf = brief_pdfs[0]
    elif final_pdf is None:
        print("__ERROR__:no PDFs came out of the converter")
        sys.exit(1)

    print(f"__PDF__:{final_pdf}")
    sys.exit(0)


if __name__ == "__main__":
    main()
