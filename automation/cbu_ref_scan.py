"""Find past LoadStar-PS (CBU) quotes and the system size each one was for.

The desk problem this solves: a new CBU enquiry for, say, 10KVA single phase is
almost always a copy of the last 10KVA single-phase quote — the Bidman lines are
identical bar the project name. Finding that last one meant remembering it.
Ask Vector cannot answer "10kva-1ph" because nothing in the mail index records a
system SIZE against a quote reference; the size only exists inside the artefacts
the sizer produces.

Two artefacts carry the size reliably, and both are ours:

  * the Tech Brief PDF (``CBU_Tech_Brief_<ref>.pdf``) that the CBU tab exports —
    a labelled page with "System size (KVA)", "Phase's", "Quote reference" and
    "Project Title" on it. A multi-system export is those pages concatenated, so
    ONE file can legitimately yield several sizes;
  * a filled copy of the LoadStar-PS sizing calculator (.xlsm/.xlsx), where
    'Sales Engineer Sheet'!B5/C17/C18 hold system / project / quote — the same
    cells cbu_export.py writes.

Anything else (PMO CBU docs, the brochure, the DualGuard-S paperwork) has no
size in it and is deliberately ignored. Mail bodies mention kVA figures too, but
a thread quoting five options lists five numbers with no way to tell which one
was quoted, so free text is not used as a source at all.

Passes, cheapest first:

  files  walk the scan directories (Downloads, Desktop, the PDF Quotes folder)
         for briefs and filled calculators. Offline, no Outlook, ~50ms a file.
  mail   optional. Uses the local mail mirror only to pick which messages carry a
         brief/calculator attachment, then opens those few through Outlook COM
         and parses the real attachment. Costs a COM round trip per message, so
         it is off unless asked for.

Usage:  python cbu_ref_scan.py --job job.json --out out.json
"""

import argparse
import datetime as dt
import json
import os
import re
import sqlite3
import sys
import tempfile
import traceback

# ── Tech Brief field labels ──────────────────────────────────────────────────
# The PDF is a spreadsheet print: every field is a label line followed by its
# value line. When a field is EMPTY the value line is simply the next label, so
# a naive "line after the label" read silently returns "Quote reference" as the
# project title (the blank template in Downloads does exactly this). Every read
# is checked against this set and dropped if it lands on another label.
BRIEF_LABELS = {
    'date generated', 'project title', 'quote reference', 'sales engineer',
    'sales engineer email', 'sales engineer phone', 'system information',
    'system size (kva)', 'capacity(w)', 'autonomy', 'ip rating', "phase's",
    'power factor', 'input voltage (ac)', 'output voltage (ac)',
    'rated output current(a)', 'heat rejection (kw)',
    'recommended input fuse (rectifier)', 'recommended input fuse (bypass)',
    'recommended minimum input cable size', 'recommended minimum output cable size',
    'maximum fault current per phase(a)', "minimum clearance's", 'rear wall',
    'top', 'front', 'side', 'general schematic',
}

BRIEF_MARKER = 'Technical Brief'

# Quote references, in the two id systems this desk runs on: Salesforce
# opportunity ids (long 006QO… and the short SR/CR/EU forms) and BidManager
# works numbers. Mirrors the patterns in dq_backfill_audit.py.
SFID18_RE     = re.compile(r'006QO0\d{4}[A-Za-z0-9]{8}')
SFID_SHORT_RE = re.compile(r'(?:SR|CR|EU)0[0O][A-Za-z0-9]{7,8}')
WORKS_RE      = re.compile(r'Q[BWVR]\d{5}', re.I)
# Short forms that lost their SR/CR/EU prefix somewhere ("00xUfRxYAK"). Only
# ever tried last, and only against a filename.
#
# NOT \b. Exported names are underscore-joined — "CBU_Tech_Brief_00xUfRxYAK_
# 1PH-10KVA.pdf" — and "_" is a word character, so \b never fires between the
# separator and the id and the whole match is lost. Same trap that cost a run
# on the D&Q audit; the boundary has to be spelled out as "not alphanumeric".
SFID_BARE_RE  = re.compile(r'(?<![A-Za-z0-9])0[0O][A-Za-z0-9]{7,10}(?![A-Za-z0-9])')

# Which attachments are worth opening in the mail pass. Name-only filter — the
# content check happens after the file is on disk.
ATT_NAME_RE = re.compile(r'(cbu.{0,4}tech.{0,4}brief|load.?star|sizing calculator|\bcbu\b)', re.I)
ATT_EXT_RE  = re.compile(r'\.(pdf|xlsm|xlsx)$', re.I)

CALC_SHEET = 'Sales Engineer Sheet'


def log(msg):
    print(msg, file=sys.stderr, flush=True)


# ── Normalising a size ───────────────────────────────────────────────────────
def norm_system(kva, phase):
    """('4', 'Single') -> '1PH- 4KVA'.  Returns (key, kva_float, phase_key)."""
    try:
        k = float(str(kva).replace(',', '.').strip())
    except (TypeError, ValueError):
        return None, None, None
    if k <= 0:
        return None, None, None
    p = str(phase or '').strip().lower()
    if p.startswith('single') or p.startswith('1'):
        ph = '1PH'
    elif p.startswith('three') or p.startswith('3'):
        ph = '3PH'
    else:
        return None, None, None
    # '%g' keeps 0.5 as "0.5" and 10.0 as "10" — the spelling cbuData.ts uses.
    return '%s- %sKVA' % (ph, '%g' % k), k, ph


def parse_system_key(text):
    """Read a key the sizer itself wrote ('3PH- 12KVA'), tolerating spacing."""
    m = re.search(r'\b([13])\s*PH\s*-?\s*([\d.]+)\s*KVA', str(text or ''), re.I)
    if not m:
        return None, None, None
    return norm_system(m.group(2), m.group(1))


def norm_ref(text):
    """Pull a quote reference out of free text — a filename or a mail subject."""
    s = str(text or '').strip()
    if not s:
        return ''
    for rx in (SFID18_RE, SFID_SHORT_RE, WORKS_RE, SFID_BARE_RE):
        m = rx.search(s)
        if m:
            ref = m.group(0)
            # Keep a revision marker if one trails the id ("CR00xxHR3YAM A2R").
            # Closed with "not alphanumeric" rather than \b for the same reason
            # as SFID_BARE_RE: in "…A1R_1PH-1KVA" the marker is followed by an
            # underscore, and \b would not see the end of it.
            rev = re.search(r'^[-\s_]*(A\d\s?R)(?![A-Za-z0-9])',
                            s[m.end():m.end() + 8], re.I)
            if rev:
                ref += ' ' + rev.group(1).upper().replace(' ', '')
            return ref
    return ''


def looks_like_ref(s):
    """True if the text carries a recognisable Salesforce or BidManager id."""
    s = str(s or '')
    return any(rx.search(s) for rx in (SFID18_RE, SFID_SHORT_RE, WORKS_RE, SFID_BARE_RE))


def field_ref(text):
    """The reference as typed into the brief's own 'Quote reference' cell.

    Taken at face value rather than run through the id patterns: this cell is
    the reference, and squeezing it through a regex loses the revision suffixes
    and prefix-less ids ("00xUfRxYAK") that the desk actually files under.
    """
    s = re.sub(r'\s+', ' ', str(text or '')).strip()
    s = re.sub(r'[-–—_\s]+$', '', s)
    if not s or len(s) > 60 or not re.search(r'[A-Za-z0-9]', s):
        return ''
    if s.lower().rstrip(':') in BRIEF_LABELS:
        return ''
    return s


def iso_date(s):
    """'27/07/2026' -> '2026-07-27'.  Anything unparseable -> ''."""
    s = str(s or '').strip()
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y'):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return ''


def mtime_date(path):
    try:
        return dt.datetime.fromtimestamp(os.path.getmtime(path)).date().isoformat()
    except OSError:
        return ''


# ── Tech Brief PDF ───────────────────────────────────────────────────────────
def brief_fields(text):
    """Label -> value for one brief page, dropping values that are labels."""
    lines = [ln.strip() for ln in text.split('\n')]
    out = {}
    for i, ln in enumerate(lines):
        key = ln.lower().rstrip(':')
        if key not in BRIEF_LABELS or key in out:
            continue
        val = lines[i + 1].strip() if i + 1 < len(lines) else ''
        if not val or val.lower().rstrip(':') in BRIEF_LABELS:
            continue
        out[key] = val
    return out


def make_rec(system, kva, phase, ref, project, duration, dated, source, detail):
    """Assemble one record, correcting the two things people get wrong.

    Reference and project get typed into each other's box often enough to be
    worth undoing: if only one of the two reads as an id, that one is the
    reference. What is left over gets a `confidence` of 'weak' — practice
    exports with "sds" in every field are in the same folders as the real work,
    and the tab dims them rather than pretending they are quotes.
    """
    ref = (ref or '').strip()
    project = (project or '').strip()
    if project and looks_like_ref(project) and not looks_like_ref(ref):
        ref, project = project, ref
    if not ref:
        return None
    return {
        'system': system, 'kva': kva, 'phase': phase,
        'quoteRef': ref, 'project': project, 'duration': duration or '',
        'dated': dated or '', 'source': source, 'detail': detail,
        'confidence': 'exact' if looks_like_ref(ref) else 'weak',
    }


def scan_pdf(path):
    """Every Tech Brief page in one PDF -> zero or more records."""
    try:
        import fitz
    except ImportError:
        raise RuntimeError('PyMuPDF (pymupdf) is not installed — pip install pymupdf')

    recs = []
    try:
        doc = fitz.open(path)
    except Exception as e:                       # not a PDF, or encrypted
        log('  skip %s: %s' % (os.path.basename(path), e))
        return recs
    try:
        for page in doc:
            try:
                text = page.get_text()
            except Exception:
                continue
            if BRIEF_MARKER not in text:
                continue
            f = brief_fields(text)
            system, kva, phase = norm_system(f.get('system size (kva)'), f.get("phase's"))
            if not system:
                continue
            # The field wins; the filename is the fallback for the ones exported
            # before the reference cell was filled in.
            ref = field_ref(f.get('quote reference')) or norm_ref(os.path.basename(path))
            rec = make_rec(system, kva, phase, ref, f.get('project title', ''),
                           f.get('autonomy', ''),
                           iso_date(f.get('date generated')) or mtime_date(path),
                           'brief', path)
            if rec:
                recs.append(rec)
    finally:
        doc.close()
    return recs


# ── Filled sizing calculator ─────────────────────────────────────────────────
def scan_calc(path):
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError('openpyxl is not installed')
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        log('  skip %s: %s' % (os.path.basename(path), e))
        return []
    try:
        if CALC_SHEET not in wb.sheetnames:
            return []
        ws = wb[CALC_SHEET]
        system, kva, phase = parse_system_key(ws['B5'].value)
        if not system:
            return []
        project = str(ws['C17'].value or '').strip()
        ref = field_ref(ws['C18'].value) or norm_ref(os.path.basename(path))
        rec = make_rec(system, kva, phase, ref, project, '',
                       mtime_date(path), 'calculator', path)
        return [rec] if rec else []
    finally:
        try:
            wb.close()
        except Exception:
            pass


def scan_file(path):
    low = path.lower()
    if low.endswith('.pdf'):
        return scan_pdf(path)
    if low.endswith('.xlsm') or low.endswith('.xlsx'):
        return scan_calc(path)
    return []


# ── Pass 1: local files ──────────────────────────────────────────────────────
# Depth 2 so "Downloads/Coventry/brief.pdf" is found but a full Desktop crawl is
# not attempted. Temporary Office lock files (~$…) are never real documents.
MAX_DEPTH = 2


def walk_dirs(dirs):
    seen = set()
    for root in dirs:
        root = os.path.expandvars(os.path.expanduser(str(root or '').strip()))
        if not root or not os.path.isdir(root):
            continue
        base_depth = root.rstrip('\\/').count(os.sep)
        for cur, subdirs, files in os.walk(root):
            if cur.count(os.sep) - base_depth >= MAX_DEPTH:
                subdirs[:] = []
            for name in files:
                if name.startswith('~$') or not ATT_EXT_RE.search(name):
                    continue
                p = os.path.join(cur, name)
                key = os.path.normcase(os.path.abspath(p))
                if key in seen:
                    continue
                seen.add(key)
                yield p


def pass_files(dirs, errors):
    recs, n = [], 0
    for path in walk_dirs(dirs):
        # A PDF is only worth opening if it could be a brief. Name-filtering
        # first keeps a 120-file Downloads folder to a handful of real reads.
        name = os.path.basename(path)
        if not ATT_NAME_RE.search(name) and not name.lower().startswith('cbu'):
            continue
        n += 1
        try:
            recs.extend(scan_file(path))
        except Exception as e:
            errors.append('%s: %s' % (name, e))
    return recs, n


# ── Pass 2: Outlook attachments ──────────────────────────────────────────────
def mail_candidates(db_path, limit=200):
    """(entry_id, store_id, subject, received) for messages carrying a brief.

    The local mirror is used only to CHOOSE the messages; their attachment names
    are indexed but their bytes are not, so the file itself still comes from
    Outlook. Filtering here is what keeps the COM pass to seconds.
    """
    if not db_path or not os.path.exists(db_path):
        return []
    con = sqlite3.connect(db_path)
    try:
        rows = con.execute(
            'SELECT entry_id, store_id, subject, received, atts FROM mail '
            'WHERE atts IS NOT NULL AND atts != "" ORDER BY received DESC').fetchall()
    finally:
        con.close()
    out = []
    for entry_id, store_id, subject, received, atts in rows:
        try:
            items = json.loads(atts or '[]')
        except (TypeError, ValueError):
            continue
        names = [i.get('name', '') for i in items if isinstance(i, dict)]
        if any(n and ATT_EXT_RE.search(n) and ATT_NAME_RE.search(n) for n in names):
            out.append((entry_id, store_id or '', subject or '', received or ''))
        if len(out) >= limit:
            break
    return out


def pass_mail(db_path, errors, limit=200):
    cands = mail_candidates(db_path, limit)
    if not cands:
        return [], 0
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    # outlook_reader falls back to automation/mail_index.db, which does not
    # exist in this layout — without this every id looks "not in index" and the
    # recovery sweep gives up on messages that are sitting right there.
    os.environ['MAGIC_MAIL_INDEX'] = os.path.abspath(db_path)
    try:
        import outlook_reader as orx
    except Exception as e:
        errors.append('Outlook reader unavailable: %s' % e)
        return [], 0
    try:
        _, ns = orx.get_outlook_ns()      # (Application, MAPI namespace)
    except Exception as e:
        errors.append('Outlook is not reachable: %s' % e)
        return [], 0

    recs, opened = [], 0
    tmp = tempfile.mkdtemp(prefix='cbu_ref_')
    for entry_id, store_id, subject, received in cands:
        try:
            item = orx.item_by_id(ns, entry_id, store_id)
        except Exception:
            item = None
        if item is None:
            continue
        opened += 1
        try:
            count = item.Attachments.Count
        except Exception:
            continue
        for k in range(1, count + 1):
            try:
                att = item.Attachments.Item(k)
                name = att.FileName or ''
            except Exception:
                continue
            if not name or not ATT_EXT_RE.search(name) or not ATT_NAME_RE.search(name):
                continue
            safe = re.sub(r'[^A-Za-z0-9._-]+', '_', name)
            dest = os.path.join(tmp, '%d_%s' % (k, safe))
            try:
                att.SaveAsFile(dest)
            except Exception as e:
                errors.append('%s: %s' % (name, e))
                continue
            try:
                found = scan_file(dest)
            except Exception as e:
                errors.append('%s: %s' % (name, e))
                found = []
            for r in found:
                # The saved copy is a scratch file; point the user at the mail.
                r['source'] = 'mail'
                r['detail'] = subject
                r['entryId'] = entry_id
                r['storeId'] = store_id
                if not r.get('dated'):
                    r['dated'] = (received or '')[:10]
                r['quoteRef'] = r['quoteRef'] or norm_ref(subject) or norm_ref(name)
                r['confidence'] = 'exact' if looks_like_ref(r['quoteRef']) else 'weak'
                if r['quoteRef']:
                    recs.append(r)
            try:
                os.remove(dest)
            except OSError:
                pass
    return recs, opened


# ── Merge ────────────────────────────────────────────────────────────────────
SOURCE_RANK = {'brief': 3, 'calculator': 2, 'mail': 1}


def _score(r):
    return (1 if r.get('confidence') == 'exact' else 0,
            SOURCE_RANK.get(r['source'], 0),
            r.get('dated') or '',
            len(r.get('project') or ''))


def dedupe(recs):
    """One row per (system, reference). Richest source and newest date win."""
    best = {}
    for r in recs:
        key = (r['system'], r['quoteRef'].upper())
        cur = best.get(key)
        if cur is None or _score(r) > _score(cur):
            best[key] = r
    out = list(best.values())
    out.sort(key=lambda r: (r['phase'], r['kva'], r.get('dated') or ''))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--job', required=True)
    ap.add_argument('--out', required=True)
    args = ap.parse_args()

    with open(args.job, encoding='utf-8') as fh:
        job = json.load(fh)

    dirs = job.get('dirs') or []
    errors = []
    result = {'ok': False, 'found': [], 'errors': errors}
    try:
        recs, files = pass_files(dirs, errors)
        log('files pass: %d candidate files, %d records' % (files, len(recs)))
        mails = 0
        if job.get('scan_mail'):
            mrecs, mails = pass_mail(job.get('mail_db') or '', errors,
                                     int(job.get('mail_limit') or 200))
            log('mail pass: %d messages opened, %d records' % (mails, len(mrecs)))
            recs.extend(mrecs)
        result.update({
            'ok': True,
            'found': dedupe(recs),
            'filesScanned': files,
            'mailsScanned': mails,
        })
    except Exception as e:
        result['error'] = str(e)
        log(traceback.format_exc())

    with open(args.out, 'w', encoding='utf-8') as fh:
        json.dump(result, fh, ensure_ascii=False, indent=1)


if __name__ == '__main__':
    main()
