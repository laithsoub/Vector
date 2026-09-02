"""
LSD Daily Work register
=======================
Every transaction the LSD tab prices gets one row in a register workbook that
mirrors Dalia's "LSD Daily work -2025 -2026.xlsx" — same header row, same order:

    Country | BU | Transaction Number | Transaction Name | Customer Number |
    Customer Name | Status | Sales Name | CPQ Last Updated | Total Value |
    Out Date | Notes | RPI Comment | RPI Comment | PV% | RPI% | RPI Value

("RPI Comment" really is there twice — the second one is kept so the column
count matches hers and a paste into her sheet lines up.)

The register is a plain .xlsx written with openpyxl, so it is readable without
Excel installed and safe to write while the pricing run has Excel busy.

Three modes, all driven by a job file the server writes:

    append   upsert one row, keyed on Transaction Number
    rows     read the register back for the tab's table
    push     post one list item per transaction to the Quotations List

    python lsd_register.py --job job.json --out result.json

The header is authoritative when the file already exists: values are written
into whatever columns the sheet already has, matched by header text. Drop
Dalia's own workbook in as the register and this writes into her columns
without reshaping them; only a register created from scratch gets the header
above.
"""

import argparse
import json
import os
import re
import sys
from datetime import datetime, date

# ─── the columns, in her order ───────────────────────────────────────────────
# (key, header). The duplicate "RPI Comment" carries a distinct key so a job can
# fill either box; the header text stays duplicated on purpose.
COLUMNS = [
    ("country",          "Country"),
    ("bu",               "BU"),
    ("transaction",      "Transaction Number"),
    ("transaction_name", "Transaction Name"),
    ("customer",         "Customer Number"),
    ("customer_name",    "Customer Name"),
    ("status",           "Status"),
    ("sales_name",       "Sales Name"),
    ("cpq_updated",      "CPQ Last Updated"),
    ("total_value",      "Total Value"),
    ("out_date",         "Out Date"),
    ("notes",            "Notes"),
    ("rpi_comment",      "RPI Comment"),
    ("rpi_comment_2",    "RPI Comment"),
    ("pv_pct",           "PV%"),
    ("rpi_pct",          "RPI%"),
    ("rpi_value",        "RPI Value"),
]
KEYS = [k for k, _ in COLUMNS]
SHEET = "LSD Daily Work"
KEY_COL = "transaction"          # what an upsert matches on

# Number formats, applied only to cells this script writes.
FORMATS = {
    "total_value": "#,##0.00",
    "rpi_value":   "#,##0.00",
    "pv_pct":      "0.0%",
    "rpi_pct":     "0.0%",
    "out_date":    "yyyy-mm-dd",
    "cpq_updated": "yyyy-mm-dd",
}
WIDTHS = {
    "country": 12, "bu": 8, "transaction": 16, "transaction_name": 34,
    "customer": 14, "customer_name": 30, "status": 16, "sales_name": 20,
    "cpq_updated": 14, "total_value": 14, "out_date": 12, "notes": 34,
    "rpi_comment": 24, "rpi_comment_2": 24, "pv_pct": 9, "rpi_pct": 9,
    "rpi_value": 14,
}
TIMEOUT = 60


def _cfg():
    """config.json next to this script — the same file every other script reads."""
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.json")
    cfg = {"sp_list": "https://eaton.sharepoint.com/sites/QuotationFactoryEMEA"}
    try:
        with open(path, encoding="utf-8") as fh:
            cfg.update(json.load(fh))
    except Exception:
        pass
    return cfg


# ─── reading and writing the workbook ────────────────────────────────────────
def _norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def _sheet(wb):
    """The register sheet. An existing workbook keeps whatever its first sheet
    is called — Dalia's is not called 'LSD Daily Work'."""
    if SHEET in wb.sheetnames:
        return wb[SHEET]
    return wb.worksheets[0]


def _header_map(ws):
    """{key: column index} from row 1. Unknown headers are ignored; a header we
    do not carry simply never gets written. The duplicate 'RPI Comment' maps to
    rpi_comment then rpi_comment_2, in the order they appear."""
    wanted = {}
    for key, head in COLUMNS:
        wanted.setdefault(_norm(head), []).append(key)
    seen, out = {}, {}
    for idx, cell in enumerate(ws[1], start=1):
        h = _norm(cell.value)
        if not h or h not in wanted:
            continue
        n = seen.get(h, 0)
        opts = wanted[h]
        if n < len(opts):
            out[opts[n]] = idx
        seen[h] = n + 1
    return out


def _create(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    for i, (key, head) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=head)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(key, 16)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)
    return wb


def _load(path, create=True):
    import openpyxl
    if not os.path.exists(path):
        if not create:
            return None
        return _create(path)
    return openpyxl.load_workbook(path)


def _coerce(key, value):
    """Job values arrive as JSON, so dates are strings and percentages may be
    either 0.062 or '6.2%'. Store real numbers/dates so the sheet sorts and
    totals like a spreadsheet, not like text."""
    if value is None or value == "":
        return None
    if key in ("total_value", "rpi_value"):
        try:
            return round(float(value), 2)
        except (TypeError, ValueError):
            return value
    if key in ("pv_pct", "rpi_pct"):
        if isinstance(value, str):
            s = value.strip().rstrip("%")
            try:
                v = float(s)
            except ValueError:
                return value
            # "6.2%" means 0.062; a bare 6.2 from a form means the same thing.
            return v / 100.0 if value.strip().endswith("%") or abs(v) > 1 else v
        try:
            return float(value)
        except (TypeError, ValueError):
            return value
    if key in ("out_date", "cpq_updated"):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        # CPQ hands back ISO timestamps; a hand-typed date is dd/mm/yyyy here.
        head = str(value).strip().replace("T", " ").split(".")[0]
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%b-%Y"):
            try:
                return datetime.strptime(head, fmt).date()
            except ValueError:
                continue
        for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(head[:10], fmt).date()
            except ValueError:
                continue
        return head
    return str(value)


def _find_row(ws, cols, transaction):
    """Row number of an existing entry for this transaction, else None."""
    col = cols.get(KEY_COL)
    if not col or not transaction:
        return None
    want = _norm(transaction)
    for r in range(2, ws.max_row + 1):
        if _norm(ws.cell(row=r, column=col).value) == want:
            return r
    return None


def _first_free_row(ws, cols):
    """Last row carrying anything in a column we know about, plus one. max_row
    over-reports on a sheet that has been edited and cleared."""
    used = [c for c in cols.values()]
    last = 1
    for r in range(2, ws.max_row + 1):
        if any(ws.cell(row=r, column=c).value not in (None, "") for c in used):
            last = r
    return last + 1


def append_row(path, row, log):
    from openpyxl.styles import Alignment
    wb = _load(path)
    ws = _sheet(wb)
    cols = _header_map(ws)
    if not cols:
        raise RuntimeError(
            f"{os.path.basename(path)} has no recognisable header row — expected "
            f"'Transaction Number', 'Total Value' and the rest on row 1.")

    txn = str(row.get(KEY_COL) or "").strip()
    at = _find_row(ws, cols, txn)
    action = "updated" if at else "appended"
    if not at:
        at = _first_free_row(ws, cols)

    written, skipped = [], []
    for key in KEYS:
        if key not in row:
            continue
        col = cols.get(key)
        if not col:
            skipped.append(key)
            continue
        val = _coerce(key, row.get(key))
        # An upsert must not blank a cell the analyst filled by hand: only a
        # value that was actually supplied overwrites.
        if val is None and action == "updated":
            continue
        c = ws.cell(row=at, column=col, value=val)
        if key in FORMATS and val is not None and not isinstance(val, str):
            c.number_format = FORMATS[key]
        if key in ("notes", "rpi_comment", "rpi_comment_2"):
            c.alignment = Alignment(wrap_text=True, vertical="top")
        written.append(key)

    try:
        wb.save(path)
    except PermissionError:
        raise RuntimeError(
            f"{os.path.basename(path)} is open in Excel — close it and register again.")
    log(f"Row {action} at line {at} of {os.path.basename(path)}"
        + (f" ({len(skipped)} column(s) not in this sheet: {', '.join(skipped)})" if skipped else ""))
    return {"action": action, "row": at, "written": written, "skipped": skipped}


def read_rows(path, limit=300):
    """The register back as JSON, newest last — the tab reverses it for display."""
    wb = _load(path, create=False)
    if wb is None:
        return {"exists": False, "rows": [], "headers": []}
    ws = _sheet(wb)
    cols = _header_map(ws)
    headers = [{"key": k, "header": h, "present": k in cols} for k, h in COLUMNS]
    out = []
    for r in range(2, ws.max_row + 1):
        rec = {}
        for key, col in cols.items():
            v = ws.cell(row=r, column=col).value
            if isinstance(v, (datetime, date)):
                v = v.strftime("%Y-%m-%d")
            rec[key] = v
        if any(v not in (None, "") for v in rec.values()):
            rec["_row"] = r
            out.append(rec)
    return {"exists": True, "rows": out[-limit:], "headers": headers,
            "total": len(out), "sheet": ws.title}


# ─── SharePoint: one list item per transaction, the road the quotes take ─────
# The register itself stays on disk. What goes to SharePoint is the DATA — one
# item per transaction in the same "Quotations List" the quote uploader writes
# to, so an LSD transaction reads like every other quote in the factory list.
#
# Automation_V4 already owns the pieces that must not drift (the INSIDE SALES
# lookup list GUID, the Salesforce-id normaliser), so they are imported rather
# than copied.
QUOTE_LIST = "Quotations List"
DEFAULT_REQUEST_TYPE = "Standard CTO"
# COUNTRY is a required choice and its options are European. A deal outside them
# posts as EMEA, with the real country kept in COMMENTS.
COUNTRY_FALLBACK = "EMEA"
COUNTRY_ALIASES = {
    "uk": "UK", "united kingdom": "UK", "gb": "UK", "great britain": "UK",
    "eire": "IRELAND", "holland": "NETHERLANDS", "czechia": "CZECH REPUBLIC",
    "deutschland": "GERMANY", "espana": "SPAIN", "italia": "ITALY",
}


def _v4():
    """Automation_V4, imported for the lookup list GUID and the SF-id rules.
    Importing is safe: everything real is behind its __main__ guard."""
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import Automation_V4 as V4
    return V4


def _session(job, host="eaton.sharepoint.com"):
    import requests
    import urllib3
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
    fed, rt = job.get("fed") or "", job.get("rt") or ""
    if not (fed and rt):
        try:
            sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
            from cookie_crypto import load_cookies
            fed, rt = load_cookies()
        except Exception as e:
            raise RuntimeError(f"Could not load the SharePoint cookies: {e}")
    if not (fed and rt):
        raise RuntimeError("Not connected to SharePoint — run Connect to JOE first.")
    s = requests.Session()
    s.cookies.set("FedAuth", fed, domain=host)
    s.cookies.set("rtFa", rt, domain=host)
    s.headers.update({"Accept": "application/json;odata=verbose"})
    s.verify = False
    return s


def _digest(session, site):
    r = session.post(f"{site}/_api/contextinfo", timeout=TIMEOUT)
    if r.status_code == 403:
        raise RuntimeError("SharePoint returned 403 — the JOE cookies have expired. "
                           "Run Connect to JOE and try again.")
    if r.status_code != 200:
        raise RuntimeError(f"SharePoint auth failed (HTTP {r.status_code}).")
    return r.json()["d"]["GetContextWebInformation"]["FormDigestValue"]


def _list_url(site, tail=""):
    return f"{site}/_api/web/lists/getbytitle('{QUOTE_LIST}'){tail}"


def _entity_type(session, site):
    r = session.get(_list_url(site, "?$select=ListItemEntityTypeFullName"), timeout=TIMEOUT)
    if r.status_code != 200:
        raise RuntimeError(f"List '{QUOTE_LIST}' could not be read (HTTP {r.status_code}).")
    return r.json()["d"]["ListItemEntityTypeFullName"]


def _choices(session, site, field):
    try:
        r = session.get(_list_url(site, f"/fields?$filter=InternalName eq '{field}'"
                                        f"&$select=InternalName,Choices"), timeout=TIMEOUT)
        res = r.json()["d"]["results"]
        return res[0]["Choices"]["results"] if res else []
    except Exception:
        return []


# ─── the extras her sheet has no column for ──────────────────────────────────
# CRM id, deal currency, the real country: the list wants them, the daily sheet
# does not carry them. They live in a sidecar JSON beside the register, keyed on
# transaction number, so the workbook itself stays exactly her layout.
def _meta_path(register):
    return os.path.join(os.path.dirname(os.path.abspath(register)), ".lsd_register_meta.json")


def _read_meta(register):
    try:
        with open(_meta_path(register), encoding="utf-8") as fh:
            return json.load(fh)
    except Exception:
        return {}


def _write_meta(register, transaction, extra):
    if not transaction or not extra:
        return
    all_meta = _read_meta(register)
    cur = all_meta.get(_norm(transaction)) or {}
    cur.update({k: v for k, v in extra.items() if v not in (None, "")})
    all_meta[_norm(transaction)] = cur
    try:
        with open(_meta_path(register), "w", encoding="utf-8") as fh:
            json.dump(all_meta, fh, indent=1, default=str)
    except Exception:
        pass


# ─── row → list item ─────────────────────────────────────────────────────────
def _country(value, allowed):
    """Her Country against the list's choices. UAE, Saudi and the rest are not in
    them, so those post as EMEA and the real country rides in COMMENTS."""
    raw = str(value or "").strip()
    if not raw:
        return COUNTRY_FALLBACK, raw
    key = _norm(raw)
    key = COUNTRY_ALIASES.get(key, key)
    for c in allowed:
        if _norm(c) == key:
            return c, raw
    return COUNTRY_FALLBACK, raw


def _division(bu):
    """FIRE-only transactions post FIRE; anything carrying EL or luminaires posts
    'EL & FIRE'. CBU/CBS has no agreed mapping yet, so it is not guessed."""
    b = _norm(bu)
    if not b:
        return None, "no BU on the row"
    if b == "fire":
        return "FIRE", ""
    if b in ("el", "luminaires", "el & fire", "el and fire", "el/fire"):
        return "EL & FIRE", ""
    if b in ("cbs", "cbu"):
        return None, "CBU/CBS has no DIVISION mapping yet"
    return None, f"BU {bu!r} does not map to a DIVISION"


def _status(value, allowed):
    """Her free-text Status onto the list's STATUS choice."""
    v = _norm(value)
    for want, choice in (("hold", "ON HOLD"), ("cancel", "CANCELED"),
                         ("escalat", "ESCALATION"), ("progress", "IN PROGRESS")):
        if want in v and choice in allowed:
            return choice
    return "PROCESSED" if "PROCESSED" in allowed else (allowed[0] if allowed else None)


def _iso(value):
    """A register date (already 'YYYY-MM-DD' or a date) as a list DateTime."""
    if isinstance(value, (datetime, date)):
        return datetime(value.year, value.month, value.day).isoformat()
    s = str(value or "").strip()[:10]
    try:
        return datetime.strptime(s, "%Y-%m-%d").isoformat()
    except ValueError:
        return datetime.now().replace(hour=0, minute=0, second=0, microsecond=0).isoformat()


def _pct(v):
    try:
        return f"{float(v) * 100:.1f}%"
    except (TypeError, ValueError):
        return ""


def _comments(row, extra, country_raw, currency):
    """Everything the list has no column for, in one readable line — the real
    country, the deal currency (PRICE is labelled EUR and nothing is converted),
    and the two RPI readings the pricing gate is judged on."""
    bits = ["LSD pricing"]
    if row.get("customer"):
        bits.append(f"customer {row['customer']}")
    if currency:
        bits.append(f"total in {currency}")
    if country_raw and _norm(country_raw) != _norm(COUNTRY_FALLBACK):
        bits.append(f"country {country_raw}")
    if row.get("pv_pct") not in (None, ""):
        bits.append(f"PV {_pct(row['pv_pct'])}")
    if row.get("rpi_pct") not in (None, ""):
        bits.append(f"RPI {_pct(row['rpi_pct'])}")
    if row.get("rpi_value") not in (None, ""):
        try:
            bits.append(f"RPI value {float(row['rpi_value']):,.2f}")
        except (TypeError, ValueError):
            pass
    for k in ("notes", "rpi_comment", "rpi_comment_2"):
        if row.get(k):
            bits.append(str(row[k]))
    if extra.get("crm") and not extra.get("sfid_ok"):
        bits.append(f"CRM {extra['crm']}")
    return " · ".join(b for b in bits if b)


def _payload(row, extra, ctx, session, digest, log):
    """One list item, or (None, why) when the row cannot be posted."""
    V4 = ctx["V4"]
    division, why = _division(row.get("bu"))
    if not division:
        return None, why

    country, country_raw = _country(row.get("country"), ctx["countries"])
    currency = extra.get("currency") or ""
    status = _status(row.get("status"), ctx["statuses"])
    day = _iso(row.get("out_date"))
    item = {
        "__metadata": {"type": ctx["entity"]},
        "Title": str(row.get("transaction") or "").strip(),
        "QUOTATION_x0020_NAME": str(row.get("transaction_name") or "").strip() or None,
        "CUSTOMER": str(row.get("customer_name") or "").strip() or None,
        "DIVISION": division,
        "Country": country,
        "REQUEST_x0020_TYPE": ctx["request_type"],
        "STATUS": status,
        # The list validates the whole date chain:
        #   ARRIVED ON <= ON-HOLD <= STARTING/RECOVERING, and when the status is
        #   Processed, STARTING/RECOVERING <= PROCESSED ON.
        # A blank middle date fails it (blank reads as 1899), which is why the
        # quote CSV carries all four. An LSD transaction is priced in one sitting
        # and has no hold or recovery, so all four are the same day — the chain
        # holds on equality.
        "ARRIVED_x0020_ON": day,
        "ON_x002d_HOLD_x0020_date": day,
        "RECOVERING_x0020_date_x0020_": day,
    }
    # ...and when it is NOT processed, PROCESSED ON must be blank — the same rule,
    # read the other way.
    if status == "PROCESSED":
        item["SENT_x0020_ON_x0020_"] = day

    # C360 ID is validated as "1-" plus 7 characters. The SAP customer number is
    # not that, so it only goes in when it really is a C360 id; otherwise it is
    # kept in COMMENTS with the rest.
    c360 = str(row.get("customer") or "").strip()
    if re.match(r"^1-.{7}$", c360):
        item["C360_x0020_ID"] = c360
    # PRICE is labelled [EUR] but the raw deal total goes in unconverted, with
    # the currency named in COMMENTS — so nothing is silently rebased at an FX
    # rate nobody sees.
    try:
        if row.get("total_value") not in (None, ""):
            item["PRICE"] = float(row["total_value"])
    except (TypeError, ValueError):
        pass

    # The list validates SALESFORCE ID as blank or exactly 18 characters, so a
    # short or odd id goes to COMMENTS rather than failing the whole item.
    sfid = V4.normalize_sfid(extra.get("crm") or "")
    if sfid and len(sfid) == 18:
        item["SALESFORCEID"] = sfid
        extra["sfid_ok"] = True
    elif sfid:
        log(f"{item['Title']}: CRM id {sfid!r} is {len(sfid)} chars, not 18 — "
            f"left out of SALESFORCEID and noted in COMMENTS instead.", "warn")

    # INSIDE SALES is a required lookup; the uploader creates the entry when the
    # person is not in the lookup list yet, exactly as the quote path does — but
    # a dry run must not write anything at all, so it only resolves and reports
    # a name that is not there yet.
    who = str(row.get("sales_name") or "").strip() or ctx["inside_sales"]
    if who:
        is_id = (V4.get_inside_sales_id(session, who) if ctx.get("dry_run")
                 else V4.ensure_inside_sales_id(session, who, digest))
        if is_id:
            item["INSIDE_x0020_SALES0Id"] = is_id
        elif ctx.get("dry_run"):
            item["INSIDE_x0020_SALES0Id"] = None
            log(f"{item['Title']}: INSIDE SALES {who!r} is not in the lookup list — "
                f"a real post would create it.", "warn")
        else:
            return None, f"INSIDE SALES {who!r} could not be resolved or created"
    else:
        return None, "no INSIDE SALES name (set one in Settings → LSD Pricing)"

    item["COMMENTS"] = _comments(row, extra, country_raw, currency)
    return {k: v for k, v in item.items() if v is not None}, ""


def _existing_id(session, site, transaction):
    """The list item already holding this transaction number, if any."""
    t = str(transaction or "").replace("'", "''")
    if not t:
        return None
    url = _list_url(site, f"/items?$select=Id&$filter=Title eq '{t}'&$top=1")
    try:
        r = session.get(url, timeout=TIMEOUT)
        res = r.json()["d"]["results"]
        return res[0]["Id"] if res else None
    except Exception:
        return None


def push(path, job, log):
    """Post register rows to the Quotations List — add, or update the item that
    already carries that transaction number."""
    cfg = _cfg()
    site = (job.get("sp_list") or cfg.get("sp_list") or "").rstrip("/")
    if not site:
        raise RuntimeError("No quotations site configured (Settings → sp_list).")
    data = read_rows(path, limit=10000)
    if not data.get("exists"):
        raise RuntimeError("There is no register yet — price and build a case first.")

    wanted = {_norm(t) for t in (job.get("transactions") or []) if str(t).strip()}
    rows = [r for r in data["rows"]
            if not wanted or _norm(r.get("transaction")) in wanted]
    if not rows:
        raise RuntimeError("No matching register rows to post.")

    session = _session(job)
    digest = _digest(session, site)
    log(f"Connected to {site}")
    ctx = {
        "V4": _v4(),
        "entity": _entity_type(session, site),
        "countries": _choices(session, site, "Country"),
        "statuses": _choices(session, site, "STATUS"),
        "request_type": job.get("request_type") or cfg.get("lsd_request_type")
                        or DEFAULT_REQUEST_TYPE,
        "inside_sales": job.get("inside_sales") or cfg.get("lsd_sales_name")
                        or cfg.get("inside_sales") or "",
        # A dry run resolves and reports; it never creates a lookup entry and
        # never touches the list.
        "dry_run": bool(job.get("dry_run")),
    }
    types = _choices(session, site, "REQUEST_x0020_TYPE")
    if types and ctx["request_type"] not in types:
        raise RuntimeError(f"REQUEST TYPE {ctx['request_type']!r} is not one of {types}.")

    meta = _read_meta(path)
    headers = {"Accept": "application/json;odata=verbose",
               "Content-Type": "application/json;odata=verbose",
               "X-RequestDigest": digest}
    added, updated, results = 0, 0, []
    for row in rows:
        txn = str(row.get("transaction") or "").strip()
        label = txn or str(row.get("transaction_name") or "?")
        extra = dict(meta.get(_norm(txn)) or {})
        try:
            item, why = _payload(row, extra, ctx, session, digest, log)
        except Exception as e:
            item, why = None, str(e)
        if not item:
            log(f"{label}: not posted — {why}", "warn")
            results.append({"transaction": txn, "action": "skipped", "reason": why})
            continue
        if job.get("dry_run"):
            results.append({"transaction": txn, "action": "dry-run",
                            "payload": {k: v for k, v in item.items() if k != "__metadata"}})
            continue
        try:
            existing = _existing_id(session, site, txn)
            if existing:
                r = session.post(_list_url(site, f"/items({existing})"), json=item,
                                 headers={**headers, "X-HTTP-Method": "MERGE", "IF-MATCH": "*"},
                                 timeout=TIMEOUT)
                ok = r.status_code in (200, 201, 204)
                if ok:
                    updated += 1
                    log(f"{label}: updated item {existing}")
                results.append({"transaction": txn, "id": existing,
                                "action": "updated" if ok else "failed",
                                "reason": "" if ok else f"HTTP {r.status_code} {r.text[:160]}"})
            else:
                r = session.post(_list_url(site, "/items"), json=item,
                                 headers=headers, timeout=TIMEOUT)
                ok = r.status_code in (200, 201)
                new_id = r.json()["d"]["Id"] if ok else None
                if ok:
                    added += 1
                    log(f"{label}: added item {new_id}")
                results.append({"transaction": txn, "id": new_id,
                                "action": "added" if ok else "failed",
                                "reason": "" if ok else f"HTTP {r.status_code} {r.text[:160]}"})
            if not ok:
                log(f"{label}: FAILED — HTTP {r.status_code} {r.text[:200]}", "error")
        except Exception as e:
            log(f"{label}: failed — {e}", "error")
            results.append({"transaction": txn, "action": "failed", "reason": str(e)})

    failed = sum(1 for x in results if x["action"] == "failed")
    skipped = sum(1 for x in results if x["action"] == "skipped")
    log(f"{added} added, {updated} updated, {skipped} skipped, {failed} failed.",
        "ok" if not failed else "warn")
    return {"site": site, "list": QUOTE_LIST, "added": added, "updated": updated,
            "skipped": skipped, "failed": failed, "results": results}



# ─── entry point ─────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--job", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

    with open(args.job, encoding="utf-8") as fh:
        job = json.load(fh)

    log_lines = []

    def log(msg, kind="info"):
        log_lines.append({"kind": kind, "msg": msg})
        print(f"[{kind}] {msg}", flush=True)

    result = {"ok": False}
    try:
        path = job.get("register")
        if not path:
            raise RuntimeError("No register path was given.")
        mode = job.get("mode") or "rows"
        if mode == "append":
            row = job.get("row") or {}
            if not str(row.get(KEY_COL) or "").strip() and not str(row.get("transaction_name") or "").strip():
                raise RuntimeError("A register row needs a transaction number or a transaction name.")
            res = append_row(path, row, log)
            # CRM id, deal currency and anything else the list wants but her
            # sheet has no column for goes to the sidecar, keyed on transaction.
            _write_meta(path, row.get(KEY_COL), job.get("extra") or {})
            result = {"ok": True, "mode": mode, "register": path, **res,
                      **read_rows(path)}
        elif mode == "rows":
            result = {"ok": True, "mode": mode, "register": path,
                      **read_rows(path, int(job.get("limit") or 300)),
                      "meta": _read_meta(path)}
        elif mode == "push":
            res = push(path, job, log)
            result = {"ok": True, "mode": mode, "register": path, **res}
        else:
            raise RuntimeError(f"Unknown mode {mode!r}.")
    except Exception as e:
        log(str(e), "error")
        result = {"ok": False, "error": str(e)}
    finally:
        result["log"] = log_lines
        with open(args.out, "w", encoding="utf-8") as fh:
            json.dump(result, fh, default=str)

    sys.exit(0 if result.get("ok") else 1)


if __name__ == "__main__":
    main()
